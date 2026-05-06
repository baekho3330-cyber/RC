"""
리뷰 수집 프로그램 - 로컬 실행용
실행: python app.py  (브라우저 자동 오픈)
"""
import json
import os
import queue
import re
import subprocess
import sys
import threading
import webbrowser
from datetime import datetime
from pathlib import Path

try:
    from flask import Flask, Response, jsonify, render_template, request, send_file
except ImportError:
    print("Flask not installed. Run: pip install flask")
    sys.exit(1)

os.chdir(Path(__file__).parent)

CONFIG_FILE = Path(__file__).parent / "config.json"
OUTPUT_DIR  = Path(__file__).parent / "output"
OUTPUT_DIR.mkdir(exist_ok=True)

SITE_PATTERNS = {
    "livart":  r"hyundailivart\.co\.kr",
    "ohou":    r"(?:store\.)?ohou\.se",
    "hanssem": r"store\.hanssem\.com",
}
SITE_NAMES = {
    "livart":  "현대리바트몰",
    "ohou":    "오늘의집",
    "hanssem": "한샘몰",
}

flask_app = Flask(__name__)
log_q: queue.Queue = queue.Queue()
state: dict = {"running": False, "result": None}


# ── 설정 ──────────────────────────────────────────────────────────────────

def load_config() -> dict:
    try:
        return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_config(data: dict):
    try:
        existing = load_config()
        existing.update(data)
        CONFIG_FILE.write_text(
            json.dumps(existing, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    except Exception:
        pass


def detect_site(text: str) -> str | None:
    for site, pat in SITE_PATTERNS.items():
        if re.search(pat, text, re.IGNORECASE):
            return site
    return None


# ── Excel 저장 ────────────────────────────────────────────────────────────

def save_to_excel(reviews: list[dict], site: str, product_name: str, output_dir: str) -> str:
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill

    Path(output_dir).mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame(reviews)
    if "내용" in df.columns and "리뷰내용" not in df.columns:
        df = df.rename(columns={"내용": "리뷰내용"})
    if "리뷰내용" in df.columns:
        before = len(df)
        df = df.drop_duplicates(subset=["리뷰내용"])
        df = df[df["리뷰내용"].str.strip().str.len() > 0]
        removed = before - len(df)
        if removed:
            print(f"[*] 중복/빈 리뷰 제거: {removed}건")
    df = df.reset_index(drop=True)

    ts   = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe = re.sub(r'[\\/:*?"<>|]', "_", product_name)[:30]
    path = f"{output_dir}/{site}_{safe}_{ts}.xlsx"

    COL_WIDTHS = {
        "리뷰ID": 14, "상품명": 40, "상품번호": 14,
        "별점": 8, "리뷰내용": 80,
        "작성자": 15, "작성일": 12, "구매옵션": 25,
        "미디어타입": 10, "도움수": 10,
        "품질점수": 10, "디자인점수": 10, "배송점수": 10, "가격점수": 10,
        "판매자답변": 40,
    }

    H_FILL = PatternFill(start_color="111111", end_color="111111", fill_type="solid")
    H_FONT = Font(color="FFFFFF", bold=True)
    content_col = next((i for i, c in enumerate(df.columns, 1) if c == "리뷰내용"), None)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="리뷰데이터", index=False)
        ws = writer.sheets["리뷰데이터"]
        for cell in ws[1]:
            cell.fill = H_FILL
            cell.font = H_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = COL_WIDTHS.get(col, 15)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    wrap_text=(cell.column == content_col), vertical="top"
                )
        for rn in range(2, ws.max_row + 1):
            ws.row_dimensions[rn].height = 55

    return path


# ── 로그 스트림 ───────────────────────────────────────────────────────────

class QueueWriter:
    def __init__(self, q: queue.Queue):
        self.q   = q
        self.buf = ""

    def write(self, text: str):
        self.buf += text
        while "\n" in self.buf:
            line, self.buf = self.buf.split("\n", 1)
            if line:
                self.q.put(line)

    def flush(self):
        if self.buf:
            self.q.put(self.buf)
            self.buf = ""


# ── 크롤링 워커 ───────────────────────────────────────────────────────────

def crawl_worker(url: str, site: str, save_dir: str, max_r: int):
    old_stdout       = sys.stdout
    sys.stdout       = QueueWriter(log_q)
    state["running"] = True
    state["result"]  = None

    try:
        print(f"[*] {SITE_NAMES[site]} 리뷰 수집 시작")

        if site == "livart":
            from crawler import crawl_livart_reviews
            reviews = crawl_livart_reviews(url, max_reviews=max_r)
        elif site == "ohou":
            from ohou_analyzer import extract_product_id, fetch_all_reviews
            pid     = extract_product_id(url)
            reviews, _ = fetch_all_reviews(pid)
            reviews = reviews[:max_r]
        elif site == "hanssem":
            from hanssem_crawler import crawl_hanssem_reviews
            reviews = crawl_hanssem_reviews(url, max_reviews=max_r)
        else:
            reviews = []

        if not reviews:
            print("[!] 수집된 리뷰가 없습니다.")
            state["result"] = {"ok": False, "msg": "리뷰 없음"}
            return

        print(f"[완료] {len(reviews)}개 리뷰 수집")

        product_name = next(
            (r.get("상품명") or r.get("상품") for r in reviews
             if r.get("상품명") or r.get("상품")), "상품"
        ) or "상품"

        print("[*] 엑셀 저장 중...")
        path = save_to_excel(reviews, site, product_name, save_dir)
        print(f"[완료] {Path(path).name}")

        state["result"] = {"ok": True, "path": path, "name": Path(path).name}

    except Exception as e:
        import traceback
        print(f"[!] 오류: {e}")
        print(traceback.format_exc())
        state["result"] = {"ok": False, "msg": str(e)}

    finally:
        sys.stdout       = old_stdout
        state["running"] = False
        log_q.put("__DONE__")


# ── Flask 라우트 ──────────────────────────────────────────────────────────

@flask_app.route("/")
def index():
    cfg = load_config()
    return render_template(
        "index.html",
        default_dir=cfg.get("save_dir", str(OUTPUT_DIR)),
        max_reviews=cfg.get("max_reviews", 500),
    )


@flask_app.route("/browse", methods=["POST"])
def browse():
    current = (request.json or {}).get("current", str(Path.home()))
    try:
        res = subprocess.run(
            ["powershell", "-Command",
             f'Add-Type -AssemblyName System.Windows.Forms; '
             f'$f = New-Object System.Windows.Forms.FolderBrowserDialog; '
             f'$f.SelectedPath = "{current}"; '
             f'$f.ShowDialog() | Out-Null; '
             f'$f.SelectedPath'],
            capture_output=True, text=True, timeout=60,
        )
        path = res.stdout.strip()
        if path:
            return jsonify({"path": path})
    except Exception:
        pass
    return jsonify({"path": ""})


@flask_app.route("/crawl", methods=["POST"])
def crawl():
    if state["running"]:
        return jsonify({"ok": False, "msg": "이미 수집 중입니다."}), 400

    data     = request.json or {}
    url      = (data.get("url") or "").strip()
    save_dir = (data.get("save_dir") or str(OUTPUT_DIR)).strip()
    max_r    = int(data.get("max_reviews") or 500)
    if max_r <= 0:
        max_r = 99999

    site = detect_site(url)
    if not site:
        return jsonify({"ok": False, "msg": "지원하지 않는 URL입니다."}), 400

    save_config({"save_dir": save_dir, "max_reviews": max_r})

    while not log_q.empty():
        try: log_q.get_nowait()
        except queue.Empty: break

    threading.Thread(
        target=crawl_worker,
        args=(url, site, save_dir, max_r),
        daemon=True,
    ).start()

    return jsonify({"ok": True, "site": SITE_NAMES[site]})


@flask_app.route("/stream")
def stream():
    def generate():
        while True:
            try:
                line = log_q.get(timeout=30)
                if line == "__DONE__":
                    result = state.get("result") or {}
                    yield f"data: __DONE__{json.dumps(result, ensure_ascii=False)}\n\n"
                    break
                yield f"data: {line}\n\n"
            except queue.Empty:
                yield "data: __PING__\n\n"

    return Response(
        generate(),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@flask_app.route("/open_folder", methods=["POST"])
def open_folder():
    path = (request.json or {}).get("path", "")
    if path:
        folder = str(Path(path).parent)
        if Path(folder).exists():
            os.startfile(folder)
    return jsonify({"ok": True})


# ── 진입점 ───────────────────────────────────────────────────────────────

def _free_port(port: int):
    try:
        result = subprocess.run(
            ["netstat", "-ano"], capture_output=True, text=True, timeout=5,
        )
        for line in result.stdout.splitlines():
            if f":{port}" in line and "LISTENING" in line:
                parts = line.split()
                pid = int(parts[-1])
                if pid != os.getpid():
                    subprocess.run(["taskkill", "/F", "/PID", str(pid)],
                                   capture_output=True)
    except Exception:
        pass


if __name__ == "__main__":
    port = 5001
    _free_port(port)
    threading.Timer(1.2, lambda: webbrowser.open(f"http://127.0.0.1:{port}")).start()
    flask_app.run(host="127.0.0.1", port=port, debug=False, use_reloader=False)
