"""
오늘의집 상품 리뷰 수집 + Claude AI 분석 보고서 생성
사용법: python ohou_analyzer.py
"""
import sys
import io
import os
import re
import time
import json
import requests
import pandas as pd
from pathlib import Path
from datetime import datetime
try:
    import anthropic
except ImportError:
    anthropic = None

if hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

# ── 설정 ──────────────────────────────────────────────────
OUTPUT_DIR = "output"
Path(OUTPUT_DIR).mkdir(exist_ok=True)


# ── 1. 리뷰 수집 ─────────────────────────────────────────
def extract_product_id(url: str) -> str | None:
    m = re.search(r"/goods/(\d+)", url)
    return m.group(1) if m else None


def _get_ohou_session(product_id: str) -> requests.Session:
    """Playwright로 실제 브라우저 쿠키/헤더를 획득해 requests Session 반환"""
    from playwright.sync_api import sync_playwright

    cookies_dict = {}
    headers_extra = {}

    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True, args=["--no-sandbox", "--disable-dev-shm-usage"])
            ctx = browser.new_context(
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
                locale="ko-KR",
            )
            page = ctx.new_page()
            page.goto(f"https://store.ohou.se/goods/{product_id}", wait_until="domcontentloaded", timeout=30000)
            page.wait_for_timeout(2000)

            for c in ctx.cookies():
                cookies_dict[c["name"]] = c["value"]

            browser.close()
        print("[*] 브라우저 세션 획득 완료")
    except Exception as e:
        print(f"[*] 브라우저 세션 실패, 기본 헤더 사용: {e}")

    session = requests.Session()
    session.headers.update({
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
        "Accept": "application/json, text/plain, */*",
        "Accept-Language": "ko-KR,ko;q=0.9",
        "Accept-Encoding": "gzip, deflate, br",
        "Origin": "https://store.ohou.se",
        "Referer": f"https://store.ohou.se/goods/{product_id}",
        "sec-ch-ua": '"Chromium";v="124", "Google Chrome";v="124"',
        "sec-ch-ua-mobile": "?0",
        "sec-ch-ua-platform": '"Windows"',
        "sec-fetch-dest": "empty",
        "sec-fetch-mode": "cors",
        "sec-fetch-site": "same-origin",
    })
    session.cookies.update(cookies_dict)
    return session


def fetch_all_reviews(product_id: str) -> tuple[list[dict], dict]:
    """오늘의집 리뷰 전체 수집. (reviews, stats) 반환"""
    session = _get_ohou_session(product_id)

    # 별점 통계
    stats = {}
    try:
        r = session.get(
            "https://store.ohou.se/api/goods/reviews/counts-for-stars",
            params={"productionId": product_id}, timeout=10
        )
        stats = r.json()
    except Exception:
        pass

    # 전체 리뷰
    all_reviews = []
    page = 1
    per = 20

    # 총 개수 먼저 파악
    r0 = session.get(
        "https://store.ohou.se/api/goods/reviews",
        params={"page": 1, "productionId": product_id, "per": 1,
                "order": "recent", "stars": "", "option": ""},
        timeout=10
    )
    if r0.status_code != 200 or not r0.text.strip():
        raise RuntimeError(f"오늘의집 API 오류 (HTTP {r0.status_code})")
    total = r0.json().get("totalCount", 0)
    print(f"[*] 총 리뷰: {total}개")

    while True:
        r = session.get(
            "https://store.ohou.se/api/goods/reviews",
            params={"page": page, "productionId": product_id, "per": per,
                    "order": "recent", "stars": "", "option": ""},
            timeout=15
        )
        if r.status_code != 200 or not r.text.strip():
            print(f"[!] 페이지 {page} 응답 오류 (HTTP {r.status_code}), 중단")
            break
        data = r.json()
        reviews = data.get("reviews", [])
        if not reviews:
            break

        for item in reviews:
            rv = item.get("review", {})
            pi = item.get("productionInformation", {})
            all_reviews.append({
                "리뷰ID": item.get("id"),
                "상품명": pi.get("name", ""),
                "브랜드": pi.get("brandName", ""),
                "옵션": pi.get("explain", ""),
                "별점": rv.get("starAvg", ""),
                "리뷰내용": rv.get("comment", ""),
                "작성일": item.get("createdAt", ""),
                "작성자": item.get("writerNickname", ""),
                "리뷰타입": item.get("reviewType", {}).get("label", ""),
                "미디어": rv.get("type", ""),
                "도움수": item.get("praiseCount", 0),
                "사용기간": next(
                    (b["name"] for b in rv.get("badges", []) if "사용" in b.get("name", "")),
                    ""
                ),
            })

        print(f"    {len(all_reviews)}/{total}개 수집")
        if len(all_reviews) >= total:
            break
        page += 1
        time.sleep(0.3)

    return all_reviews, stats


# ── 2. Excel 저장 ─────────────────────────────────────────
def save_excel(reviews: list[dict], prefix: str = "ohou") -> str:
    df = pd.DataFrame(reviews)
    df = df.drop_duplicates(subset=["리뷰내용"])
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = f"{OUTPUT_DIR}/{prefix}_reviews_{ts}.xlsx"

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="리뷰데이터", index=False)
        try:
            from openpyxl.styles import Font, PatternFill, Alignment
            ws = writer.sheets["리뷰데이터"]
            fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            for cell in ws[1]:
                cell.fill = fill
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")
            widths = {"리뷰ID": 12, "상품명": 35, "브랜드": 12, "옵션": 20,
                      "별점": 6, "리뷰내용": 70, "작성일": 12, "작성자": 15,
                      "리뷰타입": 15, "미디어": 8, "도움수": 8, "사용기간": 12}
            for i, col in enumerate(df.columns, 1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(col, 15)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top",
                                               wrap_text=(cell.column == 6))
            for rn in range(2, ws.max_row + 1):
                ws.row_dimensions[rn].height = 55
        except Exception:
            pass

    print(f"[*] Excel 저장: {path} ({len(df)}건)")
    return path


# ── 3. AI 분석 보고서 ─────────────────────────────────────
def build_review_text(reviews: list[dict], stats: dict) -> str:
    """Claude에 전달할 리뷰 텍스트 구성"""
    lines = []

    # 통계 요약
    total = len(reviews)
    avg = sum(r["별점"] for r in reviews if isinstance(r["별점"], (int, float))) / max(total, 1)
    star_dist = stats.get("countFor5", 0), stats.get("countFor4", 0), stats.get("countFor3", 0), \
                stats.get("countFor2", 0), stats.get("countFor1", 0)

    lines.append(f"[상품] {reviews[0]['상품명']} ({reviews[0]['브랜드']})")
    lines.append(f"[총 리뷰] {total}개 | 평균 별점: {avg:.2f}")
    lines.append(f"[별점 분포] 5점:{star_dist[0]} 4점:{star_dist[1]} 3점:{star_dist[2]} 2점:{star_dist[3]} 1점:{star_dist[4]}")
    lines.append("")

    for i, r in enumerate(reviews, 1):
        content = (r["리뷰내용"] or "").strip()
        if not content:
            continue
        lines.append(f"[{i}] 별점:{r['별점']} | 옵션:{r['옵션']} | 날짜:{r['작성일']} | 사용기간:{r['사용기간']}")
        lines.append(content)
        lines.append("")

    return "\n".join(lines)


def analyze_with_claude(review_text: str, product_name: str, api_key: str) -> str:
    """Claude Opus 4.6으로 리뷰 분석 보고서 생성 (스트리밍)"""
    client = anthropic.Anthropic(api_key=api_key)

    system = """당신은 소비자 리뷰 분석 전문가입니다.
주어진 상품 리뷰 데이터를 바탕으로 구조화된 분석 보고서를 한국어로 작성합니다.
보고서는 마케팅팀, 제품팀, 경영진이 즉시 활용할 수 있도록 명확하고 실용적으로 작성하세요."""

    prompt = f"""다음은 '{product_name}' 상품의 실제 구매자 리뷰 데이터입니다.

{review_text}

---

위 리뷰를 분석하여 다음 구조로 보고서를 작성해주세요:

## 1. 종합 요약
- 전반적인 고객 만족도 평가
- 핵심 키워드 (긍정/부정 각 5개)

## 2. 긍정 분석
- 가장 많이 언급된 강점 (상위 5개, 각 근거 리뷰 인용 포함)
- 구매 만족 주요 이유

## 3. 부정/개선 분석
- 불만 사항 및 개선 요청 (상위 5개, 근거 인용 포함)
- 반복적으로 나타나는 문제점

## 4. 옵션별 분석
- 인기 옵션 및 각 옵션 만족도 차이
- 옵션 선택 시 주의사항

## 5. 사용 맥락 분석
- 주요 구매 목적 및 사용 대상 (학생/직장인/가정 등)
- 사용 기간별 만족도 변화

## 6. 경쟁력 인사이트
- 가성비 관련 언급
- 타 브랜드·제품과의 비교 언급

## 7. 개선 제안 (우선순위 Top 3)
- 즉시 개선 가능한 사항
- 중장기 개선 사항

## 8. 마케팅 활용 포인트
- 광고 카피로 활용 가능한 실제 리뷰 문구 3개
- 타겟 고객군 제안
"""

    print("\n[*] Claude AI 분석 중...")
    full_response = ""

    with client.messages.stream(
        model="claude-opus-4-6",
        max_tokens=8000,
        system=[{"type": "text", "text": system, "cache_control": {"type": "ephemeral"}}],
        messages=[{"role": "user", "content": [{"type": "text", "text": prompt, "cache_control": {"type": "ephemeral"}}]}],
    ) as stream:
        for text in stream.text_stream:
            print(text, end="", flush=True)
            full_response += text

    print("\n")
    return full_response


def save_report(report: str, product_name: str, reviews: list[dict], stats: dict) -> str:
    """분석 보고서를 Excel 파일로 저장"""
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = f"{OUTPUT_DIR}/analysis_report_{ts}.xlsx"

    wb_data = {
        "항목": ["분석 보고서"],
        "내용": [report],
    }

    total = len(reviews)
    avg = sum(r["별점"] for r in reviews if isinstance(r["별점"], (int, float))) / max(total, 1)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        # 보고서 시트
        df_report = pd.DataFrame(wb_data)
        df_report.to_excel(writer, sheet_name="AI분석보고서", index=False)
        ws = writer.sheets["AI분석보고서"]
        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 120
        from openpyxl.styles import Alignment, Font, PatternFill
        ws["A1"].font = Font(bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        ws["B1"].font = Font(bold=True, color="FFFFFF")
        ws["B1"].fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        ws["B2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = max(400, len(report) // 3)

        # 통계 시트
        stats_data = {
            "항목": ["총 리뷰", "평균 별점", "5점", "4점", "3점", "2점", "1점",
                    "상품명", "브랜드"],
            "값": [total, round(avg, 2),
                  stats.get("countFor5", 0), stats.get("countFor4", 0),
                  stats.get("countFor3", 0), stats.get("countFor2", 0),
                  stats.get("countFor1", 0),
                  reviews[0]["상품명"] if reviews else "",
                  reviews[0]["브랜드"] if reviews else ""],
        }
        pd.DataFrame(stats_data).to_excel(writer, sheet_name="통계", index=False)

    print(f"[*] 보고서 저장: {path}")
    return path


# ── 메인 ─────────────────────────────────────────────────
def main():
    print("=" * 60)
    print("  오늘의집 리뷰 수집 + AI 분석 보고서")
    print("=" * 60)

    url = input("\n상품 URL 입력:\n> ").strip()
    if not url:
        url = "https://store.ohou.se/goods/968321"
        print(f"  (기본값 사용: {url})")

    product_id = extract_product_id(url)
    if not product_id:
        print("[!] 상품 ID를 추출할 수 없습니다.")
        sys.exit(1)

    api_key = os.environ.get("ANTHROPIC_API_KEY", "").strip()
    if not api_key:
        api_key = input("\nAnthropic API 키 입력 (sk-ant-...):\n> ").strip()
    if not api_key:
        print("[!] API 키가 없습니다.")
        sys.exit(1)

    print(f"\n[*] 상품 ID: {product_id}")
    print("-" * 60)

    # 1. 리뷰 수집
    reviews, stats = fetch_all_reviews(product_id)
    if not reviews:
        print("[!] 리뷰 없음")
        sys.exit(1)

    # 2. Excel 저장
    excel_path = save_excel(reviews, prefix=f"ohou_{product_id}")

    # 3. AI 분석
    review_text = build_review_text(reviews, stats)
    product_name = reviews[0]["상품명"] if reviews else "상품"

    report = analyze_with_claude(review_text, product_name, api_key)

    # 4. 보고서 저장
    report_path = save_report(report, product_name, reviews, stats)

    print("\n" + "=" * 60)
    print(f"[완료]")
    print(f"  리뷰 데이터 : {excel_path}")
    print(f"  AI 분석 보고서: {report_path}")


if __name__ == "__main__":
    main()
